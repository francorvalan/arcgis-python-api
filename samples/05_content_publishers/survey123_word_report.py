"""Utilidades para exportar registros de Survey123 a un reporte Word (.docx).

Ejemplo de uso:

    from arcgis.gis import GIS
    from survey123_word_report import (
        get_survey_layer,
        list_layer_fields,
        export_survey_to_docx,
    )

    gis = GIS("home")
    survey_item = gis.content.get("<item_id_del_survey>")
    rel_fs = survey_item.related_items("Survey2Service", "forward")[0]
    layer = get_survey_layer(rel_fs)

    # Ver columnas disponibles
    print(list_layer_fields(layer))

    # Exportar reporte con columnas seleccionadas
    export_survey_to_docx(
        layer,
        output_path="reporte_survey.docx",
        id_field="id_sitios",
        columns=["fecha", "tecnico", "estado"],
    )
"""

from __future__ import annotations

from io import BytesIO
from pathlib import Path
from tempfile import TemporaryDirectory
from typing import Iterable, Sequence


def get_survey_layer(related_feature_service, layer_index: int = 0):
    """Devuelve la capa de respuestas del Feature Service relacionado a un Survey.

    Parameters
    ----------
    related_feature_service : arcgis.gis.Item
        Item obtenido normalmente con
        `survey_item.related_items("Survey2Service", "forward")[0]`.
    layer_index : int, optional
        Índice de capa dentro de `related_feature_service.layers`.
    """

    layers = getattr(related_feature_service, "layers", None)
    if not layers:
        raise ValueError("El Feature Service relacionado no contiene capas.")

    try:
        return layers[layer_index]
    except IndexError as exc:
        raise ValueError(
            f"No existe la capa en el índice {layer_index}. "
            f"Capas disponibles: 0..{len(layers) - 1}."
        ) from exc


def list_layer_fields(layer, include_system_fields: bool = False) -> list[str]:
    """Lista nombres de campos de una capa para elegir columnas del reporte."""

    system_fields = {"OBJECTID", "GLOBALID", "SHAPE", "SHAPE_LENGTH", "SHAPE_AREA"}
    field_names = [f["name"] for f in layer.properties.fields]
    if include_system_fields:
        return field_names

    return [name for name in field_names if name.upper() not in system_fields]


def export_survey_to_docx(
    layer,
    output_path: str | Path,
    id_field: str = "id_sitios",
    columns: Sequence[str] | None = None,
    where: str = "1=1",
    photo_column_title: str = "fotos",
    image_width_inches: float = 1.3,
) -> Path:
    """Genera un documento Word con una fila por registro y fotos en la última columna.

    Parameters
    ----------
    layer : arcgis.features.FeatureLayer
        Capa de respuestas del survey.
    output_path : str | Path
        Ruta del archivo `.docx` de salida.
    id_field : str, default "id_sitios"
        Campo identificador que se usará como primera columna.
    columns : Sequence[str] | None
        Lista de campos a incluir (además de `id_field`).
        Si es `None`, incluye todos los campos no del sistema.
    where : str, default "1=1"
        Filtro SQL opcional para limitar registros.
    photo_column_title : str, default "fotos"
        Título de la última columna que incluirá adjuntos de imagen.
    image_width_inches : float, default 1.3
        Ancho de cada foto en pulgadas.
    """

    try:
        from docx import Document
        from docx.shared import Inches
    except ImportError as exc:
        raise ImportError(
            "Falta dependencia `python-docx`. Instala con: pip install python-docx"
        ) from exc

    available_fields = list_layer_fields(layer, include_system_fields=True)
    if id_field not in available_fields:
        raise ValueError(f"El campo id_field '{id_field}' no existe en la capa.")

    if columns is None:
        columns = list_layer_fields(layer)
    else:
        _validate_columns(columns, available_fields)

    # Evita duplicar id_field si ya viene en columns
    selected_columns = [id_field] + [col for col in columns if col != id_field]

    query_result = layer.query(where=where, out_fields=",".join(selected_columns))
    features = query_result.features

    document = Document()
    document.add_heading("Reporte Survey123", level=1)
    document.add_paragraph(f"Total de registros: {len(features)}")

    table = document.add_table(rows=1, cols=len(selected_columns) + 1)
    table.style = "Table Grid"

    header_cells = table.rows[0].cells
    for idx, field_name in enumerate(selected_columns):
        header_cells[idx].text = field_name
    header_cells[-1].text = photo_column_title

    oid_field = layer.properties.objectIdField

    with TemporaryDirectory() as tmp_dir:
        tmp_dir_path = Path(tmp_dir)

        for feature in features:
            row_cells = table.add_row().cells
            attrs = feature.attributes

            for idx, field_name in enumerate(selected_columns):
                value = attrs.get(field_name)
                row_cells[idx].text = "" if value is None else str(value)

            oid_value = attrs.get(oid_field)
            if oid_value is None:
                row_cells[-1].text = "Sin OBJECTID"
                continue

            image_paths = _download_image_attachments(layer, oid_value, tmp_dir_path)
            if not image_paths:
                row_cells[-1].text = "Sin fotos"
                continue

            paragraph = row_cells[-1].paragraphs[0]
            for image_path in image_paths:
                run = paragraph.add_run()
                run.add_picture(str(image_path), width=Inches(image_width_inches))
                paragraph.add_run("\n")

    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(output_path))
    return output_path


def _validate_columns(columns: Iterable[str], available_fields: Sequence[str]) -> None:
    missing = [col for col in columns if col not in available_fields]
    if missing:
        raise ValueError(
            "Las siguientes columnas no existen en la capa: "
            + ", ".join(sorted(missing))
        )


def _download_image_attachments(layer, object_id: int, temp_dir: Path) -> list[Path]:
    """Descarga adjuntos de imagen de un registro y devuelve sus rutas temporales."""

    attachments = layer.attachments.get_list(object_id=object_id)
    image_paths: list[Path] = []

    for attachment in attachments:
        content_type = (attachment.get("contentType") or "").lower()
        if not content_type.startswith("image/"):
            continue

        attachment_id = attachment["id"]
        file_name = attachment.get("name") or f"attachment_{attachment_id}.jpg"
        target_path = temp_dir / file_name

        download_result = layer.attachments.download(
            oid=object_id,
            attachment_id=attachment_id,
            save_path=str(temp_dir),
        )

        # `download` puede devolver str, Path, lista o bytes según versión/contexto.
        if isinstance(download_result, (str, Path)):
            src_path = Path(download_result)
            if src_path.exists() and src_path != target_path:
                src_path.replace(target_path)
            elif src_path.exists():
                target_path = src_path
        elif isinstance(download_result, list) and download_result:
            src_path = Path(download_result[0])
            if src_path.exists() and src_path != target_path:
                src_path.replace(target_path)
            elif src_path.exists():
                target_path = src_path
        elif isinstance(download_result, bytes):
            target_path.write_bytes(download_result)
        elif isinstance(download_result, BytesIO):
            target_path.write_bytes(download_result.getvalue())

        if target_path.exists():
            image_paths.append(target_path)

    return image_paths
